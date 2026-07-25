"""Export assembled papers as markdown / docx / pdf (with formula images)."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from exam_bank import store
from exam_bank.export_rich import (
    eq_placeholder_label,
    resolve_eq_png,
    sanitize_question_fields,
    split_eq_segments,
)

EXPORT_FORMATS = ("markdown", "docx", "pdf")

_MEDIA = {
    "markdown": ("text/markdown; charset=utf-8", ".md"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docx",
    ),
    "pdf": ("application/pdf", ".pdf"),
}


def _safe_filename(title: str, ext: str) -> str:
    base = re.sub(r'[\\/:*?"<>|]+', "_", (title or "试卷").strip()) or "试卷"
    return f"{base[:80]}{ext}"


def _ordered_questions(paper: dict[str, Any]) -> list[dict[str, Any]]:
    qids = paper.get("question_ids") or []
    out: list[dict[str, Any]] = []
    for qid in qids:
        row = store.get_question(str(qid))
        if row:
            out.append(sanitize_question_fields(row))
    return out


def _register_cjk_font() -> str | None:
    """Register a system CJK font for ReportLab; return font name or None."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return None

    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyh.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
    ]
    for path in candidates:
        if not path.is_file():
            continue
        name = "ExamCJK"
        try:
            if path.suffix.lower() == ".ttc":
                pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
            else:
                pdfmetrics.registerFont(TTFont(name, str(path)))
            return name
        except Exception:
            continue
    return None


def _esc_xml(s: str) -> str:
    return (
        (s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _png_display_size(path: Path, *, max_h: float = 16.0, max_w: float = 220.0) -> tuple[float, float]:
    """Return (width, height) in points for ReportLab / approximate DOCX."""
    try:
        from PIL import Image

        with Image.open(path) as im:
            w_px, h_px = im.size
    except Exception:
        return max_h * 2.5, max_h
    if h_px <= 0:
        return max_h * 2.5, max_h
    scale = max_h / float(h_px)
    w = w_px * scale
    h = max_h
    if w > max_w:
        scale = max_w / float(w_px)
        w = max_w
        h = h_px * scale
    return max(8.0, w), max(8.0, h)


def _rl_markup(text: str, media_ingest_id: str, *, img_h: float = 14.0) -> str:
    """Build ReportLab Paragraph XML with inline <img> for [[EQ:n]]."""
    chunks: list[str] = []
    for kind, val in split_eq_segments(text or ""):
        if kind == "text":
            chunks.append(_esc_xml(val).replace("\n", "<br/>"))
            continue
        path = resolve_eq_png(media_ingest_id, val)
        if path is None:
            chunks.append(_esc_xml(eq_placeholder_label(val)))
            continue
        src = str(path.resolve()).replace("\\", "/")
        w, h = _png_display_size(path, max_h=img_h)
        chunks.append(
            f'<img src="{src}" width="{w:.1f}" height="{h:.1f}" valign="middle"/>'
        )
    return "".join(chunks) or " "


def _docx_add_rich(paragraph: Any, text: str, media_ingest_id: str) -> None:
    from docx.shared import Pt

    for kind, val in split_eq_segments(text or ""):
        if kind == "text":
            if val:
                paragraph.add_run(val)
            continue
        path = resolve_eq_png(media_ingest_id, val)
        if path is None:
            paragraph.add_run(eq_placeholder_label(val))
            continue
        run = paragraph.add_run()
        _, h = _png_display_size(path, max_h=14.0)
        try:
            run.add_picture(str(path), height=Pt(h))
        except Exception:
            paragraph.add_run(eq_placeholder_label(val))


def _md_rich(text: str, media_ingest_id: str) -> str:
    """Markdown: replace EQ with readable label (images are binary; PDF/DOCX embed)."""
    parts: list[str] = []
    for kind, val in split_eq_segments(text or ""):
        if kind == "text":
            parts.append(val)
        else:
            path = resolve_eq_png(media_ingest_id, val)
            parts.append(eq_placeholder_label(val) if path is None else f"[公式{val}]")
    return "".join(parts)


def build_docx_bytes(*, title: str, questions: list[dict[str, Any]], include_answers: bool) -> bytes:
    """国标排版 Word：A4、2.54cm、多级列表、选项无边框表、LaTeX→可编辑 OMML。"""
    from exam_bank.export_docx_gb import build_gb_docx_bytes

    return build_gb_docx_bytes(
        title=title, questions=questions, include_answers=include_answers
    )


def build_pdf_bytes(*, title: str, questions: list[dict[str, Any]], include_answers: bool) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    from exam_bank.paper_layout import group_questions_by_qtype, section_heading

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    font = _register_cjk_font() or "Helvetica"
    title_style = ParagraphStyle(
        "ExamTitle",
        parent=styles["Heading1"],
        fontName=font,
        fontSize=16,
        leading=22,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    info_style = ParagraphStyle(
        "ExamInfo",
        parent=styles["Normal"],
        fontName=font,
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    h_style = ParagraphStyle(
        "ExamH",
        parent=styles["Heading2"],
        fontName=font,
        fontSize=12,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ExamBody",
        parent=styles["Normal"],
        fontName=font,
        fontSize=11,
        leading=18,
        spaceAfter=3,
    )

    story: list[Any] = [
        Paragraph(_esc_xml(title or "未命名试卷"), title_style),
        Paragraph(_esc_xml("姓名__________  班级__________  考号__________"), info_style),
    ]

    order, by_type = group_questions_by_qtype(questions)
    n = 1
    for si, qt in enumerate(order):
        group = by_type.get(qt) or []
        if not group:
            continue
        story.append(Paragraph(_esc_xml(section_heading(si, qt, len(group))), h_style))
        for q in group:
            mid = str(q.get("media_ingest_id") or "")
            stem = f"{n}. {q.get('stem') or ''}"
            story.append(Paragraph(_rl_markup(stem, mid), body_style))
            for opt in q.get("options") or []:
                story.append(Paragraph(_rl_markup(f"　　{opt}", mid), body_style))
            story.append(Spacer(1, 4))
            n += 1

    if include_answers:
        story.append(Paragraph(_esc_xml("参考答案与解析（考生勿看）"), h_style))
        n = 1
        for q in questions:
            mid = str(q.get("media_ingest_id") or "")
            ans = f"{n}. 【答案】{q.get('answer') or '（略）'}"
            story.append(Paragraph(_rl_markup(ans, mid), body_style))
            analysis = (q.get("analysis") or "").strip()
            if analysis:
                story.append(Paragraph(_rl_markup(f"　　【解析】{analysis}", mid), body_style))
            n += 1

    doc.build(story)
    return buf.getvalue()


def _markdown_from_questions(
    *, title: str, questions: list[dict[str, Any]], include_answers: bool
) -> str:
    from exam_bank.paper_layout import group_questions_by_qtype, section_heading

    lines = [
        f"# {title or '未命名试卷'}",
        "",
        "姓名__________  班级__________  考号__________",
        "",
    ]
    order, by_type = group_questions_by_qtype(questions)
    n = 1
    for si, qt in enumerate(order):
        group = by_type.get(qt) or []
        if not group:
            continue
        lines.append(f"## {section_heading(si, qt, len(group))}")
        lines.append("")
        for q in group:
            mid = str(q.get("media_ingest_id") or "")
            lines.append(f"{n}. {_md_rich(str(q.get('stem') or ''), mid)}")
            for opt in q.get("options") or []:
                lines.append(f"   {_md_rich(str(opt), mid)}")
            lines.append("")
            n += 1
    if include_answers:
        lines.append("## 参考答案与解析（考生勿看）")
        lines.append("")
        n = 1
        for q in questions:
            mid = str(q.get("media_ingest_id") or "")
            lines.append(f"{n}. 【答案】{_md_rich(str(q.get('answer') or '（略）'), mid)}")
            analysis = (q.get("analysis") or "").strip()
            if analysis:
                lines.append(f"   【解析】{_md_rich(analysis, mid)}")
            n += 1
    return "\n".join(lines)


def export_paper_file(
    paper: dict[str, Any],
    *,
    fmt: str,
    include_answers: bool = True,
) -> tuple[bytes, str, str]:
    """Return (bytes, content_type, filename)."""
    f = (fmt or "markdown").strip().lower()
    if f in ("md", "text"):
        f = "markdown"
    if f in ("word", "doc"):
        f = "docx"
    if f not in EXPORT_FORMATS:
        raise ValueError(f"unsupported_export_format:{fmt}")

    title = str(paper.get("title") or "未命名试卷")
    media, ext = _MEDIA[f]
    filename = _safe_filename(title, ext)
    questions = _ordered_questions(paper)

    if f == "markdown":
        # Always rebuild so EQ/noise sanitization applies (cached markdown may be stale)
        raw = _markdown_from_questions(
            title=title, questions=questions, include_answers=include_answers
        ).encode("utf-8")
        return raw, media, filename

    if f == "docx":
        return (
            build_docx_bytes(
                title=title, questions=questions, include_answers=include_answers
            ),
            media,
            filename,
        )
    return (
        build_pdf_bytes(
            title=title, questions=questions, include_answers=include_answers
        ),
        media,
        filename,
    )
