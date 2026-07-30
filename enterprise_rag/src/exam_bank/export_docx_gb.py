"""国标风格试卷 Word 导出：A4 / 2.54cm 边距 / 多级列表 / 选项无边框表 / LaTeX→OMML。

公式：题库中的 ``$...$`` / ``\\(...\\)`` / ``\\[...\\]`` 经 latex2mathml → mathml-to-omml
写入可编辑 Office Math（OMML），不是图片。``[[EQ:n]]`` 仍回退为公式图（兼容旧库）。
"""

from __future__ import annotations

import io
import logging
import re
from typing import Any

_log = logging.getLogger(__name__)

_MARGIN_CM = 2.54

# text | latex | eq
_RICH_RE = re.compile(
    r"(\$\$.+?\$\$|"
    r"\$[^$\n]+?\$|"
    r"\\\((?:.|\n)+?\\\)|"
    r"\\\[(?:.|\n)+?\\\]|"
    r"\[\[EQ:\d+\]\])",
    re.DOTALL,
)

_SUB_Q_RE = re.compile(r"(?=（\d+）)")


def split_rich_segments(raw: str) -> list[tuple[str, str]]:
    """Split into ('text'|'latex'|'eq', payload)."""
    s = raw or ""
    if not s:
        return [("text", "")]
    out: list[tuple[str, str]] = []
    last = 0
    for m in _RICH_RE.finditer(s):
        if m.start() > last:
            out.append(("text", s[last : m.start()]))
        token = m.group(1)
        if token.startswith("[[EQ:"):
            out.append(("eq", re.search(r"\d+", token).group(0)))  # type: ignore[union-attr]
        else:
            latex = token
            if latex.startswith("$$") and latex.endswith("$$"):
                latex = latex[2:-2]
            elif latex.startswith("$") and latex.endswith("$"):
                latex = latex[1:-1]
            elif latex.startswith("\\(") and latex.endswith("\\)"):
                latex = latex[2:-2]
            elif latex.startswith("\\[") and latex.endswith("\\]"):
                latex = latex[2:-2]
            out.append(("latex", latex.strip()))
        last = m.end()
    if last < len(s):
        out.append(("text", s[last:]))
    return out or [("text", s)]


def latex_to_omml_element(latex: str) -> Any | None:
    """Convert LaTeX → OMML element (``m:oMath``), or None on failure."""
    tex = (latex or "").strip()
    if not tex:
        return None
    try:
        import latex2mathml.converter as latex_conv
        import math_ml2omml
        from docx.oxml import parse_xml
    except Exception as e:  # noqa: BLE001
        _log.warning("OMML deps missing: %s", e)
        return None
    try:
        mathml = latex_conv.convert(tex)
        omml = math_ml2omml.convert(mathml)
        # ensure single root with namespace
        omml = (omml or "").strip()
        if not omml:
            return None
        if not omml.startswith("<"):
            return None
        ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        if "xmlns:m=" not in omml[:80]:
            # fragment like <m:oMath>...</m:oMath>
            wrapped = f'<w:span xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:m="{ns}">{omml}</w:span>'
            root = parse_xml(wrapped)
            # first child should be oMath
            for child in root:
                if "oMath" in child.tag:
                    return child
            return root[0] if len(root) else None
        return parse_xml(omml)
    except Exception as e:  # noqa: BLE001
        _log.warning("latex→OMML failed (%s): %s", tex[:40], e)
        return None


def _set_cell_border_none(cell: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _add_numbering_part(doc: Any) -> int:
    """Inject multilevel numbering; return numId for exam lists (levels 0/1/2)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    # Ensure numbering part exists (python-docx creates it when a list style is used)
    try:
        numbering_part = doc.part.numbering_part
    except Exception:
        tmp = doc.add_paragraph()
        tmp.style = "List Number"
        parent = tmp._element.getparent()
        parent.remove(tmp._element)
        numbering_part = doc.part.numbering_part

    numbering = numbering_part._element  # CT_Numbering

    abstract_id = 100
    num_id = 100

    for abs_el in list(numbering.findall(qn("w:abstractNum"))):
        if abs_el.get(qn("w:abstractNumId")) == str(abstract_id):
            numbering.remove(abs_el)
    for num_el in list(numbering.findall(qn("w:num"))):
        if num_el.get(qn("w:numId")) == str(num_id):
            numbering.remove(num_el)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "hybridMultilevel")
    abstract.append(multi)

    abstract.append(
        _lvl_xml(
            ilvl=0,
            start=1,
            num_fmt="chineseCounting",
            lvl_text="%1、",
            indent_cm=0.0,
            hanging_cm=0.74,
        )
    )
    abstract.append(
        _lvl_xml(
            ilvl=1,
            start=1,
            num_fmt="decimal",
            lvl_text="%2.",
            indent_cm=0.74,
            hanging_cm=0.74,
        )
    )
    abstract.append(
        _lvl_xml(
            ilvl=2,
            start=1,
            num_fmt="decimal",
            lvl_text="（%3）",
            indent_cm=1.48,
            hanging_cm=0.74,
        )
    )
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def _lvl_xml(
    *,
    ilvl: int,
    start: int,
    num_fmt: str,
    lvl_text: str,
    indent_cm: float,
    hanging_cm: float,
) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), str(ilvl))
    st = OxmlElement("w:start")
    st.set(qn("w:val"), str(start))
    lvl.append(st)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    lvl.append(fmt)
    lt = OxmlElement("w:lvlText")
    lt.set(qn("w:val"), lvl_text)
    lvl.append(lt)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    pPr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(int(Cm(indent_cm + hanging_cm))))
    ind.set(qn("w:hanging"), str(int(Cm(hanging_cm))))
    pPr.append(ind)
    lvl.append(pPr)
    return lvl


def _apply_num_pr(paragraph: Any, *, num_id: int, ilvl: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    # remove old numPr
    for old in pPr.findall(qn("w:numPr")):
        pPr.remove(old)
    numPr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numId_el = OxmlElement("w:numId")
    numId_el.set(qn("w:val"), str(num_id))
    numPr.append(ilvl_el)
    numPr.append(numId_el)
    pPr.append(numPr)


def _append_rich_to_paragraph(paragraph: Any, text: str, media_ingest_id: str = "") -> None:
    from exam_bank.export_rich import eq_placeholder_label, resolve_eq_png
    from docx.shared import Pt

    for kind, val in split_rich_segments(text or ""):
        if kind == "text":
            if val:
                paragraph.add_run(val)
            continue
        if kind == "latex":
            el = latex_to_omml_element(val)
            if el is not None:
                paragraph._p.append(el)
            else:
                paragraph.add_run(f"${val}$")
            continue
        # eq image fallback
        path = resolve_eq_png(media_ingest_id, val)
        if path is not None:
            run = paragraph.add_run()
            try:
                run.add_picture(str(path), height=Pt(14))
            except Exception:
                paragraph.add_run(eq_placeholder_label(val))
        else:
            paragraph.add_run(eq_placeholder_label(val))


def _add_options_table(doc: Any, options: list[str], media_ingest_id: str) -> None:
    """2×2（或一行）无边框表格对齐 A/B/C/D。"""
    opts = [str(o).strip() for o in options if str(o).strip()]
    if not opts:
        return
    # pad to even for 2-col
    n = len(opts)
    cols = 2 if n >= 2 else 1
    rows = (n + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.autofit = True
    idx = 0
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            _set_cell_border_none(cell)
            cell.text = ""
            if idx < n:
                p = cell.paragraphs[0]
                # clear default run
                for run in list(p.runs):
                    run._r.getparent().remove(run._r)
                _append_rich_to_paragraph(p, opts[idx], media_ingest_id)
                idx += 1


def _split_stem_subquestions(stem: str) -> tuple[str, list[str]]:
    """Return (main_stem, [sub_parts starting with （n）...])."""
    s = (stem or "").strip()
    if "（1）" not in s and "(1)" not in s:
        return s, []
    # normalize (1) → （1）
    s2 = re.sub(r"\((\d+)\)", r"（\1）", s)
    parts = [p.strip() for p in _SUB_Q_RE.split(s2) if p and p.strip()]
    if len(parts) <= 1:
        return s, []
    main = parts[0]
    subs = parts[1:]
    # if main itself starts with （1）, all are subs
    if main.startswith("（") and re.match(r"^（\d+）", main):
        return "", parts
    return main, subs


def build_gb_docx_bytes(
    *,
    title: str,
    questions: list[dict[str, Any]],
    include_answers: bool = True,
) -> bytes:
    """Build A4 / 2.54cm / multilevel list / borderless options / OMML formulas."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Mm, Pt

    from exam_bank.paper_layout import group_questions_by_qtype, section_heading
    from exam_bank.subject_catalog import qtype_label

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(_MARGIN_CM)
    section.bottom_margin = Cm(_MARGIN_CM)
    section.left_margin = Cm(_MARGIN_CM)
    section.right_margin = Cm(_MARGIN_CM)

    num_id = _add_numbering_part(doc)

    t = doc.add_paragraph()
    run = t.add_run(title or "未命名试卷")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "宋体"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph("姓名__________  班级__________  考号__________")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER

    order, by_type = group_questions_by_qtype(questions)
    for si, qt in enumerate(order):
        group = by_type.get(qt) or []
        if not group:
            continue
        # Level 0: 一、选择题… — use list + also set visible heading text without manual 一
        h = doc.add_paragraph()
        _apply_num_pr(h, num_id=num_id, ilvl=0)
        # lvlText already has 一、; body = 题型（本大题共 n 小题）
        label = qtype_label(qt)
        hr = h.add_run(f"{label}（本大题共 {len(group)} 小题）")
        hr.bold = True
        hr.font.size = Pt(12)

        for q in group:
            mid = str(q.get("media_ingest_id") or "")
            main, subs = _split_stem_subquestions(str(q.get("stem") or ""))
            p = doc.add_paragraph()
            _apply_num_pr(p, num_id=num_id, ilvl=1)
            if main:
                _append_rich_to_paragraph(p, main, mid)
            elif not subs:
                _append_rich_to_paragraph(p, str(q.get("stem") or ""), mid)

            for sub in subs:
                # strip leading （n） because list level provides it? 
                # Our lvl_text is （%3） auto — strip manual number from content
                body = re.sub(r"^（\d+）\s*", "", sub).strip()
                sp = doc.add_paragraph()
                _apply_num_pr(sp, num_id=num_id, ilvl=2)
                _append_rich_to_paragraph(sp, body, mid)

            opts = [str(o).strip() for o in (q.get("options") or []) if str(o).strip()]
            if opts:
                _add_options_table(doc, opts, mid)

    if include_answers:
        sep = doc.add_paragraph()
        sep.add_run("参考答案与解析（考生勿看）").bold = True
        n = 1
        for q in questions:
            mid = str(q.get("media_ingest_id") or "")
            p = doc.add_paragraph()
            p.add_run(f"{n}. 【答案】")
            _append_rich_to_paragraph(p, str(q.get("answer") or "（略）"), mid)
            analysis = (q.get("analysis") or "").strip()
            if analysis:
                ap = doc.add_paragraph()
                ap.add_run("　　【解析】")
                _append_rich_to_paragraph(ap, analysis, mid)
            n += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
